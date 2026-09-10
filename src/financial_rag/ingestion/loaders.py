import re
from dataclasses import dataclass
from pathlib import Path
from typing import List, Optional  # noqa: UP035

import fitz
import pandas as pd
from docx import Document
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph


@dataclass
class DocumentRecord:
    document: str
    file_type: str
    page: Optional[int] = None
    sheet: Optional[str] = None
    row_start: Optional[int] = None
    row_end: Optional[int] = None
    content_type: str = "text"
    text: str = ""
    section: Optional[str] = None
    image_path: Optional[str] = None
    source_path: Optional[str] = None
    table_id: Optional[str] = None


def clean_text(text: str) -> str:
    if not text:
        return ""

    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)

    return text.strip()


def extract_pdf_tables(
    page,
    document_name: str,
    page_number: int
) -> List[DocumentRecord]:

    records = []

    try:
        table_finder = page.find_tables()
        tables = table_finder.tables

        for table_index, table in enumerate(tables, start=1):

            extracted = table.extract()

            if not extracted:
                continue

            rows = []

            for row in extracted:
                cleaned_row = [
                    clean_text(str(cell)) if cell is not None else ""
                    for cell in row
                ]

                if any(cleaned_row):
                    rows.append(cleaned_row)

            if not rows:
                continue

            headers = rows[0]

            headers = [
                header if header else f"Column {i + 1}"
                for i, header in enumerate(headers)
            ]

            output = []

            for row_index, row in enumerate(rows[1:], start=1):

                row_parts = []

                for column_index, value in enumerate(row):

                    if column_index < len(headers):
                        column_name = headers[column_index]
                    else:
                        column_name = f"Column {column_index + 1}"

                    if value:
                        row_parts.append(
                            f"{column_name}: {value}"
                        )

                if row_parts:
                    output.append(
                        f"Row {row_index}: "
                        + " | ".join(row_parts)
                    )

            if not output:
                continue

            table_text = "\n".join(output)

            records.append(
                DocumentRecord(
                    document=document_name,
                    file_type="pdf",
                    page=page_number,
                    content_type="table",
                    text=(
                        f"Table {table_index}\n\n"
                        f"{table_text}"
                    ),
                    table_id=(
                        f"{document_name}_"
                        f"table_{table_index}"
                    ),
                )
            )

    except Exception as e:
        print(
            f"Warning: failed to extract PDF tables "
            f"from {document_name}, page {page_number}: {e}"
        )

    return records


def extract_pdf_text(
    page,
    document_name: str,
    page_number: int
) -> List[DocumentRecord]:

    records = []

    try:
        table_finder = page.find_tables()
        tables = table_finder.tables

        table_rects = [
            table.bbox
            for table in tables
        ]

    except Exception:
        table_rects = []

    try:
        blocks = page.get_text("blocks")

        text_parts = []

        for block in blocks:
            x0, y0, x1, y1, text = block[:5]

            if not text or not text.strip():
                continue

            block_rect = fitz.Rect(
                x0,
                y0,
                x1,
                y1
            )

            overlaps_table = any(
                block_rect.intersects(
                    fitz.Rect(table_rect)
                )
                for table_rect in table_rects
            )

            if overlaps_table:
                continue

            cleaned = clean_text(text)

            if cleaned:
                text_parts.append(cleaned)

        page_text = "\n\n".join(text_parts)

        if page_text:
            records.append(
                DocumentRecord(
                    document=document_name,
                    file_type="pdf",
                    page=page_number,
                    content_type="text",
                    text=page_text,
                )
            )

    except Exception as e:
        print(
            f"Warning: failed to extract text from "
            f"{document_name}, page {page_number}: {e}"
        )

    return records


def load_pdf(pdf_path: Path) -> List[DocumentRecord]:

    records = []

    try:
        doc = fitz.open(pdf_path)

        if doc.needs_pass:
            raise ValueError(
                f"PDF is password protected: {pdf_path.name}"
            )

        document_name = pdf_path.name

        for page_index, page in enumerate(doc):

            page_number = page_index + 1

            text_records = extract_pdf_text(
                page=page,
                document_name=document_name,
                page_number=page_number,
            )

            records.extend(text_records)

            table_records = extract_pdf_tables(
                page=page,
                document_name=document_name,
                page_number=page_number,
            )

            records.extend(table_records)

        doc.close()

    except Exception as e:
        print(
            f"Error loading PDF "
            f"{pdf_path.name}: {e}"
        )

    return records


def iter_docx_blocks(document):

    body = document.element.body

    for child in body.iterchildren():

        if isinstance(child, CT_P):
            yield Paragraph(child, document)

        elif isinstance(child, CT_Tbl):
            yield Table(child, document)


def docx_table_to_text(table: Table) -> str:

    rows = table.rows

    if not rows:
        return ""

    output = []

    headers = [
        clean_text(cell.text)
        for cell in rows[0].cells
    ]

    headers = [
        header if header else f"Column {i + 1}"
        for i, header in enumerate(headers)
    ]

    for row_index, row in enumerate(
        rows[1:],
        start=1
    ):

        values = [
            clean_text(cell.text)
            for cell in row.cells
        ]

        row_parts = []

        for column_index, value in enumerate(values):

            if column_index < len(headers):
                column_name = headers[column_index]
            else:
                column_name = f"Column {column_index + 1}"

            if value:
                row_parts.append(
                    f"{column_name}: {value}"
                )

        if row_parts:
            output.append(
                f"Row {row_index}: "
                + " | ".join(row_parts)
            )

    return "\n".join(output)


def load_docx(docx_path: Path) -> List[DocumentRecord]:

    records = []

    try:
        document = Document(docx_path)

        document_name = docx_path.name

        heading_path = []
        table_index = 0

        for block in iter_docx_blocks(document):

            if isinstance(block, Paragraph):

                text = clean_text(block.text)

                if not text:
                    continue

                try:
                    style_name = block.style.name or ""
                except Exception:
                    style_name = ""

                if style_name.lower().startswith("heading"):

                    match = re.search(
                        r"(\d+)",
                        style_name
                    )

                    if match:
                        level = int(match.group(1))

                        heading_path = (
                            heading_path[:level - 1]
                        )

                        heading_path.append(text)

                    continue

                section = " > ".join(
                    heading_path
                )

                contextual_text = text

                if section:
                    contextual_text = (
                        f"Section: {section}\n\n"
                        f"{text}"
                    )

                records.append(
                    DocumentRecord(
                        document=document_name,
                        file_type="docx",
                        content_type="text",
                        text=contextual_text,
                        section=section or None,
                        source_path=str(docx_path),
                    )
                )

            elif isinstance(block, Table):

                table_index += 1

                table_text = docx_table_to_text(block)

                if not table_text:
                    continue

                section = " > ".join(
                    heading_path
                )

                contextual_text = (
                    f"Table {table_index}"
                )

                if section:
                    contextual_text += (
                        f"\nSection: {section}"
                    )

                contextual_text += (
                    f"\n\n{table_text}"
                )

                records.append(
                    DocumentRecord(
                        document=document_name,
                        file_type="docx",
                        content_type="table",
                        text=contextual_text,
                        section=section or None,
                        table_id=(
                            f"{document_name}_"
                            f"table_{table_index}"
                        ),
                        source_path=str(docx_path),
                    )
                )

    except Exception as e:
        print(
            f"Error loading DOCX "
            f"{docx_path.name}: {e}"
        )

    return records


def load_txt(txt_path: Path) -> List[DocumentRecord]:

    records = []

    try:

        try:
            text = txt_path.read_text(
                encoding="utf-8"
            )

        except UnicodeDecodeError:
            text = txt_path.read_text(
                encoding="latin-1"
            )

        text = clean_text(text)

        if text:
            records.append(
                DocumentRecord(
                    document=txt_path.name,
                    file_type="txt",
                    content_type="text",
                    text=text,
                    source_path=str(txt_path),
                )
            )

    except Exception as e:
        print(
            f"Error loading TXT "
            f"{txt_path.name}: {e}"
        )

    return records


def load_markdown(
    md_path: Path
) -> List[DocumentRecord]:

    records = []

    try:

        try:
            text = md_path.read_text(
                encoding="utf-8"
            )

        except UnicodeDecodeError:
            text = md_path.read_text(
                encoding="latin-1"
            )

        lines = text.splitlines()

        heading_path = []
        current_block = []

        def flush_block():

            if not current_block:
                return

            block_text = clean_text(
                "\n".join(current_block)
            )

            if not block_text:
                return

            section = " > ".join(
                heading_path
            )

            contextual_text = block_text

            if section:
                contextual_text = (
                    f"Section: {section}\n\n"
                    f"{block_text}"
                )

            records.append(
                DocumentRecord(
                    document=md_path.name,
                    file_type="markdown",
                    content_type="text",
                    text=contextual_text,
                    section=section or None,
                    source_path=str(md_path),
                )
            )

            current_block.clear()

        for line in lines:

            heading_match = re.match(
                r"^(#{1,6})\s+(.+?)\s*$",
                line
            )

            if heading_match:

                flush_block()

                level = len(
                    heading_match.group(1)
                )

                heading = clean_text(
                    heading_match.group(2)
                )

                heading_path = (
                    heading_path[:level - 1]
                )

                heading_path.append(heading)

            elif line.strip() == "":
                flush_block()

            else:
                current_block.append(line)

        flush_block()

    except Exception as e:
        print(
            f"Error loading Markdown "
            f"{md_path.name}: {e}"
        )

    return records


def format_structured_row(
    row: pd.Series,
    columns: List[str]
) -> str:

    parts = []

    for column in columns:

        value = row.get(column, "")

        if pd.isna(value):
            value = ""

        value = str(value).strip()

        if value:
            parts.append(
                f"{column}: {value}"
            )

    return "\n".join(parts)


def load_csv(
    csv_path: Path
) -> List[DocumentRecord]:

    records = []

    try:

        df = pd.read_csv(
            csv_path,
            dtype=object
        )

        df = df.fillna("")

        columns = [
            str(column).strip()
            for column in df.columns
        ]

        schema_text = (
            f"CSV file: {csv_path.name}\n\n"
            f"Columns: "
            + ", ".join(columns)
        )

        records.append(
            DocumentRecord(
                document=csv_path.name,
                file_type="csv",
                content_type="schema",
                text=schema_text,
                row_start=1,
                row_end=1,
                source_path=str(csv_path),
            )
        )

        for index, row in df.iterrows():

            row_text = format_structured_row(
                row,
                columns
            )

            if not row_text:
                continue

            physical_row = index + 2

            records.append(
                DocumentRecord(
                    document=csv_path.name,
                    file_type="csv",
                    content_type="structured_row",
                    text=row_text,
                    row_start=physical_row,
                    row_end=physical_row,
                    source_path=str(csv_path),
                )
            )

    except Exception as e:
        print(
            f"Error loading CSV "
            f"{csv_path.name}: {e}"
        )

    return records


def load_excel(
    excel_path: Path
) -> List[DocumentRecord]:

    records = []

    try:

        sheets = pd.read_excel(
            excel_path,
            sheet_name=None,
            dtype=object
        )

        for sheet_name, df in sheets.items():

            if df.empty:
                continue

            df = df.fillna("")

            columns = [
                str(column).strip()
                for column in df.columns
            ]

            schema_text = (
                f"Workbook: {excel_path.name}\n"
                f"Sheet: {sheet_name}\n\n"
                f"Columns: "
                + ", ".join(columns)
            )

            records.append(
                DocumentRecord(
                    document=excel_path.name,
                    file_type="excel",
                    sheet=str(sheet_name),
                    content_type="schema",
                    text=schema_text,
                    row_start=1,
                    row_end=1,
                    source_path=str(excel_path),
                )
            )

            for index, row in df.iterrows():

                row_text = format_structured_row(
                    row,
                    columns
                )

                if not row_text:
                    continue

                physical_row = index + 2

                records.append(
                    DocumentRecord(
                        document=excel_path.name,
                        file_type="excel",
                        sheet=str(sheet_name),
                        content_type="structured_row",
                        text=row_text,
                        row_start=physical_row,
                        row_end=physical_row,
                        source_path=str(excel_path),
                    )
                )

    except Exception as e:
        print(
            f"Error loading Excel "
            f"{excel_path.name}: {e}"
        )

    return records


LOADER_REGISTRY = {
    ".pdf": load_pdf,
    ".docx": load_docx,
    ".txt": load_txt,
    ".md": load_markdown,
    ".markdown": load_markdown,
    ".csv": load_csv,
    ".xlsx": load_excel,
    ".xls": load_excel,
}


def load_document(
    file_path: Path
) -> List[DocumentRecord]:

    file_path = Path(file_path)

    if not file_path.exists():
        raise FileNotFoundError(
            f"File not found: {file_path}"
        )

    suffix = file_path.suffix.lower()

    if suffix not in LOADER_REGISTRY:
        raise ValueError(
            f"Unsupported file type: {suffix}"
        )

    loader = LOADER_REGISTRY[suffix]

    return loader(file_path)